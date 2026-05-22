import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import Config

def export_tasks_to_excel(tasks, output_filename: str = None) -> str:
    """Exports tasks to an Excel workbook (.xlsx file).
    
    If output_filename is not specified, a default with timestamp is created in EXPORT_DIR.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Задачи"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Headers in Russian
    headers = [
        "ID", 
        "Название", 
        "Описание", 
        "Автор", 
        "Исполнитель", 
        "Дата создания", 
        "Срок выполнения", 
        "Статус", 
        "Этап (В работе)", 
        "Дата завершения",
        "Приоритет"
    ]
    
    # Style definitions
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Elegant dark blue
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    border_thin = Side(border_style="thin", color="D9D9D9")
    data_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    data_font = Font(name="Calibri", size=11)
    
    # Write headers
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = Border(left=Side(border_style="thin", color="1F4E79"),
                             right=Side(border_style="thin", color="1F4E79"))
                             
    ws.row_dimensions[1].height = 25
    
    # Map raw statuses to readable Russian values
    status_map = {
        "new": "Новая",
        "in_progress": "В работе",
        "completed": "Завершено"
    }
    
    priority_map = {
        "low": "Низкий",
        "medium": "Средний",
        "high": "Высокий"
    }

    # Write data
    for row_idx, task in enumerate(tasks, start=2):
        creator_name = task.creator.full_name if task.creator else "Неизвестно"
        assignee_name = task.assignee.full_name if task.assignee else "Не назначено"
        
        created_str = task.created_at.strftime("%Y-%m-%d %H:%M:%S") if task.created_at else ""
        completed_str = task.completed_at.strftime("%Y-%m-%d %H:%M:%S") if task.completed_at else ""
        due_str = task.due_date if task.due_date else ""
        
        status_readable = status_map.get(task.status, task.status)
        priority_readable = priority_map.get(task.priority, task.priority)
        sub_stage_name = task.sub_stage.name if (task.status == "in_progress" and task.sub_stage) else ""
        
        row_data = [
            task.id,
            task.title,
            task.description or "",
            creator_name,
            assignee_name,
            created_str,
            due_str,
            status_readable,
            sub_stage_name,
            completed_str,
            priority_readable
        ]
        
        ws.append(row_data)
        
        # Apply data cell styles
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = data_border
            # Alignment rules
            if col_idx in (1, 6, 7, 8, 9, 10, 11): # ID, dates, status, priorities
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
        ws.row_dimensions[row_idx].height = 20
        
    # Auto-adjust column widths based on longest values
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        # Cap some column widths to look organized
        width = min(max(max_len + 3, 10), 40)
        ws.column_dimensions[col_letter].width = width
        
    if not output_filename:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = os.path.join(Config.EXPORT_DIR, f"kanban_export_{timestamp}.xlsx")
        
    wb.save(output_filename)
    return output_filename
